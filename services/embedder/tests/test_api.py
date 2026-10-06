import io
import math

import fitz  # pymupdf
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app import app

client = TestClient(app)


def test_health_reports_model_and_dim():
    r = client.get("/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok"
    assert body["dim"] == 384
    assert body["model"] == "intfloat/multilingual-e5-small"


def test_embed_returns_normalized_384_vectors():
    r = client.post("/embed", json={"texts": ["Політика відпусток", "Vacation policy"], "kind": "passage"})
    assert r.status_code == 200
    body = r.json()
    assert body["dim"] == 384
    assert len(body["vectors"]) == 2
    for v in body["vectors"]:
        assert len(v) == 384
        assert abs(math.sqrt(sum(x * x for x in v)) - 1.0) < 1e-3


def test_query_and_passage_embeddings_differ_for_same_text():
    q = client.post("/embed", json={"texts": ["відпустка"], "kind": "query"}).json()["vectors"][0]
    p = client.post("/embed", json={"texts": ["відпустка"], "kind": "passage"}).json()["vectors"][0]
    assert q != p


def test_embed_rejects_empty_list_and_bad_kind():
    assert client.post("/embed", json={"texts": [], "kind": "query"}).status_code == 422
    assert client.post("/embed", json={"texts": ["a"], "kind": "other"}).status_code == 422


def _pdf_bytes(pages_text):
    doc = fitz.open()
    for text in pages_text:
        page = doc.new_page()
        # Built-in Helvetica has no Cyrillic glyphs; use PyMuPDF's bundled fallback font.
        page.insert_font(fontname="cyr", fontbuffer=fitz.Font("cjk").buffer)
        page.insert_text((72, 72), text, fontname="cyr")
    return doc.tobytes()


def test_extract_text_from_pdf_returns_pages():
    pdf = _pdf_bytes(["Перша сторінка про відпустки", "Second page about onboarding"])
    r = client.post("/extract-text", files={"file": ("doc.pdf", pdf, "application/pdf")})
    assert r.status_code == 200
    body = r.json()
    assert body["meta"]["pages_count"] == 2
    assert "відпустки" in body["pages"][0]["text"]
    assert body["pages"][1]["page"] == 2


def test_extract_text_from_docx_returns_single_page():
    d = DocxDocument()
    d.add_paragraph("Онбординг")
    d.add_paragraph("Перший день")
    buf = io.BytesIO()
    d.save(buf)
    r = client.post(
        "/extract-text",
        files={"file": ("doc.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["meta"]["pages_count"] == 1
    assert body["pages"][0]["text"] == "Онбординг\n\nПерший день"


def test_extract_text_rejects_png():
    r = client.post("/extract-text", files={"file": ("x.png", b"\x89PNG\r\n", "image/png")})
    assert r.status_code == 422
